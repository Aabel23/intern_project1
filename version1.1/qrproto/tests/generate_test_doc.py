"""One-off script: generates qrproto_test_documentation.xlsx, a detailed
row-per-test breakdown of the qrproto test suite (input, expected
output, and why each test passes or fails). Not part of the test
suite itself -- run manually whenever test_qrproto.py changes:

    ./.venv/bin/python3 qrproto/tests/generate_test_doc.py
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

OUT_PATH = Path(__file__).resolve().parent / "qrproto_test_documentation.xlsx"
RUN_PREFIX = './.venv/bin/python3 -m pytest "qrproto/tests/test_qrproto.py::'

# Each row: (group, node_id, description, input, expected_output, why_pass_fail)
ROWS = [
    # --- CRC -----------------------------------------------------------
    ("CRC (crc.py)", "test_crc_known_vector",
     "Confirms the CRC-16/CCITT-FALSE implementation itself is correct, "
     "independent of anything else in the protocol, against the "
     "algorithm's own published test vector.",
     'crc16_ccitt(b"123456789")',
     "0x29B1 (the standard CRC-16/CCITT-FALSE check value)",
     "PASS if the function returns 0x29B1. FAIL means poly (0x1021), "
     "init value (0xFFFF), or the no-reflection/no-final-XOR rules were "
     "implemented wrong -- every OCRC/ICRC in the system would then be "
     "silently incompatible with any external tool computing the same "
     "checksum by hand."),

    ("CRC (crc.py)", "test_compute_crc_is_five_digits",
     "Confirms compute_crc() always renders as a fixed 5-digit "
     "zero-padded decimal string, never shorter.",
     'compute_crc("004200")',
     "A string of length 5 (e.g. \"01978\")",
     "PASS if len(result) == 5. FAIL would mean a low CRC value renders "
     "without leading zeros (e.g. \"178\" instead of \"00178\"), which "
     "would silently shift every field after it in the payload."),

    # --- BCD -------------------------------------------------------------
    ("BCD packing (envelope.py)", "test_bcd_round_trip_odd_length",
     "Confirms bcd_pack -> bcd_unpack recovers the original digit "
     "string. Every real inner payload is 27+8n digits, always odd, so "
     "this uses an odd-length input to match real usage.",
     'bcd_pack("123") then bcd_unpack() on the result',
     '"123" (the original string, unchanged)',
     "PASS if the round trip returns the exact input. FAIL means the "
     "0xF pad-nibble convention is broken, which would corrupt every "
     "payload's plaintext before it even reaches AES-GCM."),

    ("BCD packing (envelope.py)", "test_bcd_unpack_rejects_wrong_pad",
     "Confirms bcd_unpack() rejects a plaintext whose final nibble is "
     "not the 0xF sentinel, instead of silently treating it as a real "
     "digit.",
     "bcd_pack(\"123\") with its last nibble corrupted from 0xF to 0x3",
     "FormatError raised",
     "PASS if FormatError is raised. FAIL (no exception, or wrong "
     "exception) would mean a corrupted or truncated decryption result "
     "could silently produce one extra fake digit instead of being "
     "caught."),

    # --- inner layer -------------------------------------------------------
    *[
        ("Inner layer (inner.py)", f"test_inner_round_trip_every_count[{n}]",
         f"Round-trips a payload with exactly {n} instruction pair(s) "
         "through encode_inner() then decode_inner(), one test per "
         "opcode count 0 through 12 (13 tests total).",
         f"sku=42, {n} Instruction objects (alternating Grams/bool), "
         "timestamp=1700000000, machine_id=123",
         f"Inner digit string of length {27 + 8*n}; decoded sku=42, "
         f"count={n}, timestamp and machine_id unchanged, ingredients "
         "in original order",
         "PASS if every field survives the round trip and the length "
         "matches 27+8n exactly. FAIL at any single count would mean "
         "the field-offset math (opcode_offset/data_offset/ts_offset/"
         "mid_offset) is wrong for that specific pair count.")
        for n in range(13)
    ],

    ("Inner layer (inner.py)", "test_inner_mid_reserved_rejected_at_decode",
     "Confirms MID = 000000 (reserved, spec section 3.6) is rejected "
     "when decoding, even though it is structurally well-formed.",
     "A hand-built inner body: SKU=0042, count=00, TS=1700000000, "
     "MID=000000, correct CRC",
     "ProtocolError raised, message contains \"reserved\"",
     "PASS if the reserved MID is caught. FAIL would let an "
     "uninitialised-variable or zeroed-buffer bug produce a payload "
     "that looks valid but targets no real machine."),

    ("Inner layer (inner.py)", "test_inner_duplicate_ingredient_rejected_at_encode",
     "Confirms encode_inner() refuses to build a payload naming the "
     "same ingredient twice -- caught at encode time, not just decode.",
     "Two Instruction objects, both ingredient=7 (one True, one "
     "Grams(5))",
     "ProtocolError raised, message contains \"duplicate\"",
     "PASS if encoding is refused before a payload is produced. FAIL "
     "would mean a bad label could physically be printed."),

    ("Inner layer (inner.py)", "test_inner_ingredient_out_of_range[0]",
     "Confirms ingredient 00 (below the valid 01-24 range) is rejected.",
     "Instruction(ingredient=0, value=True)",
     "ProtocolError raised at encode_inner()",
     "PASS if 00 is refused. FAIL would let a null/uninitialised "
     "ingredient id slip into a real payload."),

    ("Inner layer (inner.py)", "test_inner_ingredient_out_of_range[25]",
     "Confirms ingredient 25 (above the valid 01-24 range) is rejected.",
     "Instruction(ingredient=25, value=True)",
     "ProtocolError raised at encode_inner()",
     "PASS if 25 is refused. FAIL would let a typo'd or future "
     "not-yet-assigned ingredient id through."),

    ("Inner layer (inner.py)", "test_weight_boundary_values_round_trip[0]",
     "Confirms Grams(0) -- the lower boundary of the weight field -- "
     "survives an encode/decode round trip intact.",
     "Instruction(ingredient=1, value=Grams(0))",
     "Decoded instruction value == 0",
     "PASS if 0 g round-trips exactly. FAIL would mean the lower "
     "boundary is mishandled (e.g. treated as falsy and dropped)."),

    ("Inner layer (inner.py)", "test_weight_boundary_values_round_trip[9999]",
     "Confirms Grams(9999) -- the upper boundary (the 4-digit field's "
     "ceiling) -- survives an encode/decode round trip intact.",
     "Instruction(ingredient=1, value=Grams(9999))",
     "Decoded instruction value == 9999",
     "PASS if 9999 g round-trips exactly. FAIL would mean the ceiling "
     "value overflows the 4-digit field or is off by one."),

    ("Inner layer (inner.py)", "test_weight_above_boundary_rejected",
     "Confirms a weight one gram past the ceiling is refused at "
     "encode time.",
     "Instruction(ingredient=1, value=Grams(10000)).encode()",
     "ProtocolError raised",
     "PASS if 10000 g is refused. FAIL would silently truncate or "
     "wrap the value into a 4-digit field, pouring the wrong amount."),

    ("Inner layer (inner.py)", "test_boolean_invalid_data_rejected",
     "Confirms a boolean-type pair with data other than 0000/0001 is "
     "rejected at decode.",
     "A hand-crafted pair: ingredient=07, type=02 (boolean), data=0007",
     "ProtocolError raised, message contains \"boolean\"",
     "PASS if data=0007 is refused. FAIL would let an arbitrary number "
     "be silently interpreted as a yes/no flag."),

    ("Inner layer (inner.py)", "test_registry_membership_enforced",
     "Confirms decode_inner(..., registry=...) rejects an ingredient "
     "not present in the supplied registry set, and accepts one that "
     "is.",
     "A payload with ingredient=7; decoded once with registry={1,2,3} "
     "and once with registry={7}",
     "First call raises ProtocolError (\"registry\"); second call "
     "succeeds",
     "PASS if the negative and positive registry checks both behave "
     "correctly. FAIL on either half means the registry gate is not "
     "actually enforced (or is too strict)."),

    # --- envelope ----------------------------------------------------------
    ("Envelope encrypt/decrypt (envelope.py)", "test_envelope_round_trip",
     "Confirms encrypt_payload() -> decrypt_payload() recovers the "
     "exact original inner digit string, and the outer string starts "
     "with the VER marker \"14\".",
     "An inner payload for sku=42 with 2 instructions, encrypted with "
     "a fixed 32-byte test key",
     "Decrypted inner digits == original inner digits; outer[:2] == "
     '"14"',
     "PASS if both checks hold. FAIL would mean the AES-GCM/BCD/"
     "decimal-rendering pipeline is not actually reversible."),

    ("Envelope encrypt/decrypt (envelope.py)", "test_envelope_wrong_key_is_crypto_error",
     "Confirms decrypting with a different key raises CryptoError "
     "specifically, not a generic or unhandled exception.",
     "A valid encrypted payload; decrypt_payload() called with every "
     "key byte incremented by 1",
     "CryptoError raised (AES-GCM authentication tag fails)",
     "PASS if CryptoError is raised. FAIL (wrong exception type, or "
     "no exception) would mean a wrong-key scan is not cleanly "
     "reported as a cryptographic failure."),

    ("Envelope encrypt/decrypt (envelope.py)", "test_envelope_bad_ocrc_is_checksum_error",
     "Confirms flipping the last printed digit (inside the OCRC "
     "trailer) is caught by the outer CRC check, before decryption is "
     "ever attempted.",
     "A valid payload with its last character changed",
     "ChecksumError raised",
     "PASS if ChecksumError is raised (not CryptoError). FAIL or the "
     "wrong exception type would mean label damage and cryptographic "
     "tampering can no longer be told apart, which is the entire "
     "point of having two separate CRCs."),

    ("Envelope encrypt/decrypt (envelope.py)", "test_envelope_bad_ver_is_format_error",
     "Confirms changing VER from \"14\" to \"13\" is rejected at stage "
     "1, before the CRC or the AES key are even consulted.",
     "A valid payload with its first two characters changed to \"13\"",
     "FormatError raised",
     "PASS if FormatError is raised. FAIL would mean a stale or "
     "foreign-version payload could reach decryption instead of being "
     "rejected on sight."),

    ("Envelope encrypt/decrypt (envelope.py)", "test_envelope_wrong_length_is_format_error",
     "Confirms a digit string whose length is not one of the 13 valid "
     "outer lengths (for counts 0-12) is rejected outright.",
     '"140" repeated 20 times (60 digits -- not in the length table)',
     "FormatError raised",
     "PASS if the length is rejected immediately. FAIL would mean an "
     "arbitrary-length string could reach further processing and "
     "potentially crash on a bad slice."),

    ("Envelope encrypt/decrypt (envelope.py)", "test_old_v13_length_never_collides",
     "Confirms the plaintext v1.3 protocol's possible lengths (11+8m) "
     "never intersect this protocol's outer lengths (for any count "
     "0-12) -- the property that lets a v1.4 scanner reject a stale "
     "label on length alone, no decryption needed.",
     "The full set of v1.3 lengths {11+8m : m=0..12} and this "
     "protocol's outer lengths {outer_len(n) : n=0..12}",
     "Empty intersection (set())",
     "PASS if the two length sets share no value. FAIL would mean an "
     "old plaintext v1.3 label could be mistaken for a v1.4 payload of "
     "the same length, which the design assumes can never happen."),

    ("Envelope encrypt/decrypt (envelope.py)", "test_envelope_value_overflow_is_format_error",
     "Confirms a hand-crafted payload with a technically-correct OCRC "
     "but an envelope value too large to fit its byte count (>= "
     "256^B) is caught as a structural error, not a raw Python "
     "OverflowError crash.",
     "A 109-digit payload (count=0 shape) with the envelope digits set "
     "to exactly 256^42 and a valid OCRC over it",
     'FormatError raised, message contains "exceeds"',
     "PASS if FormatError is raised cleanly. FAIL would mean a "
     "hand-crafted malicious payload (buildable without the key, since "
     "OCRC needs no secret) could crash the decoder instead of being "
     "rejected."),

    # --- api: full 8-stage verify -------------------------------------------
    ("Full verification (api.py)", "test_create_verify_happy_path",
     "End-to-end confirmation that create() -> verify() preserves the "
     "SKU and every instruction under normal, valid conditions.",
     "sku=42, 2 instructions (Grams(13), True), machine_id=123, "
     "now=1700000000, verified with the same machine_id and now",
     "result[\"sku\"] == 42; instructions == [(7, 13), (12, True)]",
     "PASS if the full pipeline round-trips correctly end to end. FAIL "
     "here would mean something in create()/verify() itself is broken, "
     "even if every lower-level module passes on its own."),

    ("Full verification (api.py)", "test_stage6_machine_mismatch",
     "Confirms verify() rejects a payload when the scanner's "
     "machine_id differs from the one baked into the payload.",
     "A payload created for machine_id=123, verified with "
     "machine_id=124",
     "MachineMismatchError raised",
     "PASS if the mismatch is caught. FAIL would mean any scanner "
     "could redeem a label meant for a different machine -- the exact "
     "misdirection risk MID exists to close."),

    ("Full verification (api.py)", "test_stage7_expiry_boundary_is_inclusive[600-True]",
     "Confirms a payload verified exactly at max_age_seconds (600s "
     "old) still passes -- the boundary is inclusive.",
     "A payload created at now=T, verified at now=T+600",
     "verify() returns normally (no exception)",
     "PASS if no exception is raised at exactly 600s. FAIL (an "
     "exception here) would mean the boundary is off by one, rejecting "
     "valid labels prematurely."),

    ("Full verification (api.py)", "test_stage7_expiry_boundary_is_inclusive[601-False]",
     "Confirms a payload verified one second past max_age_seconds "
     "(601s old) is rejected.",
     "A payload created at now=T, verified at now=T+601",
     "ExpiredError raised",
     "PASS if ExpiredError is raised at 601s. FAIL (no exception) "
     "would mean expired labels never actually expire."),

    ("Full verification (api.py)", "test_stage7_clock_skew_boundary_is_inclusive[30-True]",
     "Confirms a payload timestamped exactly clock_skew_seconds (30s) "
     "ahead of the scanner's clock still passes.",
     "A payload created at now=T+30, verified at now=T",
     "verify() returns normally (no exception)",
     "PASS if no exception is raised at exactly 30s ahead. FAIL would "
     "mean normal NTP clock drift between generator and scanner "
     "rejects otherwise-valid labels."),

    ("Full verification (api.py)", "test_stage7_clock_skew_boundary_is_inclusive[31-False]",
     "Confirms a payload timestamped 31s ahead of the scanner's clock "
     "(one second past the allowance) is rejected.",
     "A payload created at now=T+31, verified at now=T",
     "ClockSkewError raised",
     "PASS if ClockSkewError is raised at 31s ahead. FAIL (no "
     "exception) would mean a payload probing or exploiting clock "
     "skew is never caught."),

    ("Full verification (api.py)", "test_stage_ordering_expiry_wins_over_semantics",
     "Confirms a payload that is BOTH expired AND fails a stage-8 "
     "semantic check (ingredient not in registry) is reported as "
     "expired -- stage 7 must gate stage 8, per the spec's mandated "
     "order.",
     "A payload with ingredient=7, verified 601s late with "
     "registry={99} (7 is not in it either)",
     "ExpiredError raised (not a registry/ProtocolError)",
     "PASS if ExpiredError specifically is raised. FAIL (a different "
     "exception type) would mean an operator sees a misleading "
     "\"unknown ingredient\" message instead of the true, earlier-"
     "gating problem (an expired label)."),

    ("Full verification (api.py)", "test_stage_ordering_machine_mismatch_wins_over_expiry",
     "Confirms a payload that is BOTH mistargeted AND expired is "
     "reported as a machine mismatch -- stage 6 must gate stage 7.",
     "A payload created for machine_id=123, verified with "
     "machine_id=124 AND 601s late",
     "MachineMismatchError raised (not ExpiredError)",
     "PASS if MachineMismatchError specifically is raised. FAIL would "
     "mean the wrong diagnosis reaches the operator (\"reprint, it's "
     "expired\" instead of \"walk it to the right machine\")."),

    ("Full verification (api.py)", "test_validate_key_rejects_wrong_length",
     "Confirms a key shorter than 32 bytes is rejected immediately "
     "with a clear ValueError.",
     '_validate_key(b"short")',
     "ValueError raised",
     "PASS if the short key is caught up front. FAIL would mean a "
     "misconfigured key crashes deep inside the AES-GCM library "
     "instead of failing with a readable message."),

    ("Full verification (api.py)", "test_validate_key_rejects_wrong_type",
     "Confirms a non-bytes key (e.g. a str) is rejected the same way.",
     '_validate_key("not bytes")',
     "ValueError raised",
     "PASS if the wrong type is caught. FAIL would mean a caller "
     "passing a str key by mistake gets a confusing low-level error "
     "instead of a clear one."),

    ("Full verification (api.py)", "test_normalize_machine_id_rejects_float",
     "Confirms machine_id=123.7 is refused rather than silently "
     "truncated to 123 -- stage 6 is a security check, so a sloppy "
     "caller must fail closed.",
     "_normalize_machine_id(123.7)",
     "ValueError raised",
     "PASS if the float is refused outright. FAIL (silent truncation "
     "to 123) would let a caller's bug accidentally match a payload "
     "against a machine that was never actually named."),

    ("Full verification (api.py)", "test_normalize_machine_id_accepts_digit_string",
     "Confirms machine_id passed as a digit string (as it might arrive "
     "from a config file) is accepted and normalized to an int.",
     '_normalize_machine_id("000123")',
     "123 (int)",
     "PASS if the string normalizes correctly, including leading "
     "zeros. FAIL would break machine_id values loaded from text-based "
     "config sources."),

    # --- qr sizing ---------------------------------------------------------
    ("QR sizing (qr.py)", "test_version_table_matches_spec_5_2",
     "Confirms the auto-derived VERSION_FOR_COUNT table matches the "
     "spec's QR version table exactly for every opcode count.",
     "VERSION_FOR_COUNT (built at import time from constants.outer_len)",
     "{0:6, 1:6, 2:6, 3:6, 4:7, 5:8, 6:8, 7:8, 8:8, 9:8, 10:9, 11:9, "
     "12:9}",
     "PASS if every count maps to the spec's version. FAIL would mean "
     "a printed label could be sized for a QR version too small to "
     "hold it, producing an unscannable symbol."),

    ("QR sizing (qr.py)", "test_minimum_version_matches_payload_count",
     "Confirms minimum_version() correctly reports the QR version for "
     "a real 5-instruction payload.",
     "A payload with 5 instructions (ingredients 1-5, all boolean "
     "True)",
     "minimum_version(payload) == 8",
     "PASS if it reports version 8, matching the table above for "
     "count=5. FAIL would mean the printed label is sized wrong."),

    ("QR sizing (qr.py)", "test_symbol_dimensions_version_6",
     "Confirms the physical size formula for QR version 6 at 0.5mm "
     "module size matches the spec's worked figures.",
     "symbol_dimensions(6)",
     "modules=41, symbol_mm=20.5, footprint_mm=24.5",
     "PASS if all three figures match. FAIL would mean the label "
     "footprint quoted to a label vendor is wrong."),

    # --- package surface -----------------------------------------------
    ("Package surface (__init__.py)", "test_top_level_reexports_work",
     "Confirms ChecksumError imported from the top-level qrproto "
     "package (not qrproto.errors) is the exact same exception class "
     "that decrypt_payload() actually raises -- i.e. the __init__.py "
     "re-export is wired to the real class, not a duplicate.",
     "A payload with its last digit flipped, decrypted, caught with "
     "\"except ChecksumError\" where ChecksumError was imported from "
     "qrproto (top level)",
     "The top-level ChecksumError catches the exception successfully",
     "PASS if the top-level import catches it. FAIL would mean "
     "\"from qrproto import ChecksumError\" is a different class "
     "object than the one raised internally, breaking any external "
     "code that catches errors via the top-level import."),
]


FLOW_PREFIX = './.venv/bin/python3 -m pytest "qrproto/tests/test_flow.py::TestOrderFlow::'

# One realistic order (SKU 15, sugar dial -> grams, tapioca boolean,
# machine_id=1) walking through the whole pipeline in order. Each row:
# (step, node_id, description, input, expected_output, why_pass_fail, depends_on)
FLOW_ROWS = [
    (1, "test_01_store_builds_the_order",
     "The customer's choices become opcode/data pairs. Mirrors "
     "store_gui/drinks-pos.js: the sugar dial (a UI concept, 20%) is "
     "converted to a final weight in grams client-side BEFORE it "
     "becomes a pair -- the protocol itself never sees a percentage.",
     "sugar_baseline_g=50, customer_sugar_percent=20",
     "sugar_grams == 10; order_instructions = "
     "[Instruction(7, Grams(10)), Instruction(12, True)] (saved for "
     "later steps)",
     "PASS if the gram conversion is exactly 10 and the instruction "
     "list is built. FAIL here means the very first assumption of the "
     "story (what the customer ordered) is already wrong, so every "
     "later step would be checking the wrong thing.",
     "none (first step)"),

    (2, "test_02_server_creates_the_payload",
     "The server turns the order from step 1 into an encrypted, "
     "signed, machine-bound QR payload addressed to line 1.",
     "sku=15, instructions from step 1, machine_id=1, "
     "key=DEPLOYMENT_KEY, now=1700000000 (ORDER_PLACED_AT)",
     "A 128-digit outer payload string starting with \"14\" (saved "
     "for later steps)",
     "PASS if create() returns a payload without raising. FAIL means "
     "api.create() itself is broken -- an order that was correctly "
     "built in step 1 cannot even be turned into a printable label.",
     "step 1 (order_instructions)"),

    (3, "test_03_payload_shape_is_valid",
     "Sanity-checks the printed string before it goes anywhere near a "
     "printer: correct VER marker, and a length the QR version table "
     "actually accounts for.",
     "The payload string from step 2",
     'payload[:2] == "14"; minimum_version(payload) == 6',
     "PASS if both checks hold. FAIL would mean step 2 produced a "
     "string that looks right but is not actually a well-formed v1.4 "
     "payload -- caught here before wasting a print.",
     "step 2 (payload)"),

    (4, "test_04_label_prints_as_qr",
     "Renders the payload as an actual QR image, exactly as the label "
     "printer would.",
     "The payload string from step 2",
     "A 294x294 px PIL image, saved to a temp file "
     "(qrproto_flow_label.png)",
     "PASS if make_qr() accepts the payload and produces an image. "
     "FAIL would mean a payload that passed every structural check "
     "still cannot physically become a scannable label (e.g. capacity "
     "exceeded, or a missing image dependency).",
     "step 2 (payload)"),

    (5, "test_05_scanner_reads_and_verifies_at_line_1",
     "The customer walks the label to line 1 a minute later. The "
     "scanner (configured with machine_id=1) reads the digits and "
     "verifies the payload.",
     "The payload from step 2; key=DEPLOYMENT_KEY, machine_id=1, "
     "now=ORDER_PLACED_AT+60",
     "verify() returns a result dict (saved for step 6); no exception",
     "PASS if the correct machine, correct key, and a fresh timestamp "
     "are accepted. FAIL here means the label that was just printed "
     "cannot be redeemed under completely normal conditions -- the "
     "core happy path is broken.",
     "step 2 (payload)"),

    (6, "test_06_pour_instructions_match_the_order",
     "What the machine is about to pour must be exactly what the "
     "customer ordered in step 1, no more, no less.",
     "verified_result from step 5; order_instructions from step 1",
     "poured == ordered == [(7, 10), (12, True)]; sku == 15",
     "PASS if the decoded instructions are byte-for-byte the same as "
     "what was ordered. FAIL would mean data is being silently altered "
     "somewhere in the encrypt/decrypt round trip -- the drink poured "
     "would not match what the customer paid for.",
     "step 1 (order) + step 5 (verified_result)"),

    (7, "test_07_same_label_rejected_at_line_2",
     "The exact same physical label is shown to the WRONG machine "
     "(line 2 instead of line 1).",
     "The payload from step 2; machine_id=2 (LINE_2_MACHINE_ID), "
     "now=ORDER_PLACED_AT+60",
     "MachineMismatchError raised",
     "PASS if the wrong machine is rejected. FAIL would mean MID "
     "binding does not actually work, and a label printed for one line "
     "could be redeemed on any other line in the shop.",
     "step 2 (payload)"),

    (8, "test_08_same_label_rejected_after_closing",
     "The same label is found on the floor and scanned the next "
     "morning -- 12 hours later, well past the 600s default freshness "
     "window.",
     "The payload from step 2; now=ORDER_PLACED_AT+43200 (12h later)",
     "ExpiredError raised",
     "PASS if a stale label is rejected. FAIL would mean printed "
     "labels never actually expire, so any discarded or lost label "
     "could be redeemed indefinitely.",
     "step 2 (payload)"),

    (9, "test_09_scratched_label_reported_as_damage_not_forgery",
     "The label got scuffed in a pocket: one printed digit misread. "
     "This must fail the OUTER CRC (label damage -- \"reprint it\"), "
     "not reach decryption and fail as a crypto error (which would "
     "read as \"someone tampered with this\").",
     "The payload from step 2 with its last character flipped",
     "ChecksumError raised (specifically, not CryptoError)",
     "PASS if ChecksumError specifically is raised. FAIL (wrong "
     "exception, e.g. CryptoError) would mean an operator sees "
     '"possible tampering" for an ordinary scuffed label instead of '
     '"reprint it" -- the wrong diagnosis for a routine problem.',
     "step 2 (payload)"),
]


def build_run_command(prefix: str, node_id: str) -> str:
    return prefix + node_id + '"'


def write_sheet(ws, headers, table_rows, widths):
    ws.append(headers)

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    wrap = Alignment(vertical="top", wrap_text=True)
    for i, row in enumerate(table_rows, start=1):
        ws.append(row)
        for col in range(1, len(headers) + 1):
            ws.cell(row=i + 1, column=col).alignment = wrap

    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(table_rows) + 1}"


def main() -> None:
    wb = Workbook()

    # --- sheet 1: unit tests (test_qrproto.py) --------------------------
    ws1 = wb.active
    ws1.title = "Unit tests (test_qrproto)"
    headers1 = ["#", "Group", "Test Name (pytest node id)", "Run Command",
                "Description", "Input", "Expected Output", "Why It Passes / Fails"]
    rows1 = [
        [i, group, node_id, build_run_command(RUN_PREFIX, node_id), desc, inp, out, why]
        for i, (group, node_id, desc, inp, out, why) in enumerate(ROWS, start=1)
    ]
    write_sheet(ws1, headers1, rows1, widths=[4, 28, 46, 62, 46, 40, 40, 60])

    # --- sheet 2: flow test (test_flow.py) ------------------------------
    ws2 = wb.create_sheet("Flow test (test_flow)")
    headers2 = ["Step", "Test Name (pytest node id)", "Run Command", "Depends On",
                "Description", "Input", "Expected Output", "Why It Passes / Fails"]
    rows2 = [
        [step, node_id, build_run_command(FLOW_PREFIX, node_id), depends,
         desc, inp, out, why]
        for step, node_id, desc, inp, out, why, depends in FLOW_ROWS
    ]
    write_sheet(ws2, headers2, rows2, widths=[6, 46, 78, 30, 46, 40, 40, 60])

    wb.save(OUT_PATH)
    print(f"wrote {len(ROWS)} unit-test rows + {len(FLOW_ROWS)} flow-test rows to {OUT_PATH}")


if __name__ == "__main__":
    main()
